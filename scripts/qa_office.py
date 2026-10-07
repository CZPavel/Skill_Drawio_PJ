"""Optional Office target probe. Creates new files only; rendering is optional."""
import argparse
import importlib
import json
from pathlib import Path
import shutil
import subprocess
import tempfile


def qa_office(image, output_dir, target='document', render=False):
    image = Path(image).resolve()
    if target not in ('document', 'presentation'):
        raise ValueError('target must be document or presentation')
    if not image.is_file():
        raise FileNotFoundError(image)
    if image.suffix.lower() not in ('.png', '.jpg', '.jpeg'):
        raise ValueError('Use a PNG/JPEG export for portable Office insertion')
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    artifact = output_dir / ('diagram-probe.docx' if target == 'document' else 'diagram-probe.pptx')
    pdf = artifact.with_suffix('.pdf')
    if artifact.exists() or pdf.exists():
        raise FileExistsError('Use a fresh output directory; existing probe files are preserved')
    report = {'target': target, 'artifact': None, 'rendered_pdf': None,
              'status': 'unverified', 'visual_review_required': True}
    try:
        module = importlib.import_module('docx' if target == 'document' else 'pptx')
    except ImportError:
        report['reason'] = 'Optional python-docx/python-pptx package unavailable; no Office target test performed'
        return report
    if target == 'document':
        from docx.shared import Mm
        document = module.Document()
        section = document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(20)
        document.add_picture(str(image), width=Mm(170))
        document.save(artifact)
    else:
        from pptx.util import Inches
        from PIL import Image
        presentation = module.Presentation()
        presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        with Image.open(image) as picture:
            ratio = picture.width / picture.height
        width = min(12.333, 6.5 * ratio)
        height = width / ratio
        slide.shapes.add_picture(str(image), Inches((13.333-width)/2), Inches((7.5-height)/2), width=Inches(width), height=Inches(height))
        presentation.save(artifact)
    report['artifact'] = str(artifact)
    report['reason'] = 'Office file created; rendering not requested'
    if render:
        executable = shutil.which('soffice') or shutil.which('libreoffice')
        if not executable:
            report['reason'] = 'LibreOffice unavailable; file insertion verified, target rendering unverified'
        else:
            # Isolated user profile avoids touching a running LibreOffice session.
            with tempfile.TemporaryDirectory() as profile:
                command = [executable, '-env:UserInstallation=' + Path(profile).as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(output_dir.resolve()), str(artifact.resolve())]
                try:
                    result = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
                    if result.returncode == 0 and pdf.is_file() and pdf.stat().st_size > 0:
                        report.update(status='rendered', rendered_pdf=str(pdf), reason='PDF created; visual review still required')
                    else:
                        report['reason'] = 'LibreOffice conversion failed: ' + (result.stderr or result.stdout).strip()
                except (OSError, subprocess.TimeoutExpired) as error:
                    report['reason'] = 'LibreOffice conversion failed: ' + str(error)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('image', type=Path)
    parser.add_argument('output_dir', type=Path)
    parser.add_argument('--target', choices=('document', 'presentation'), default='document')
    parser.add_argument('--render', action='store_true')
    args = parser.parse_args()
    print(json.dumps(qa_office(args.image, args.output_dir, args.target, args.render), indent=2))


if __name__ == '__main__':
    main()
